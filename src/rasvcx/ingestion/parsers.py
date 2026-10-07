"""M17 document parsers — bounded-memory extraction in an isolated subprocess.

Public API: run_parser_subprocess(file_path, fmt_or_filename, settings=None)
The 2nd arg may be a format ("txt") or a filename ("sample.txt"); the format is
derived from the extension when needed. `settings` is optional.
"""

from __future__ import annotations

import csv
import logging
import multiprocessing
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace
from typing import Any

logger = logging.getLogger(__name__)

_KNOWN_FMTS = {"pdf", "txt", "md", "html", "csv", "xlsx", "docx", "pptx", "json", "xml", "zip"}
_EXT_TO_FMT = {
    ".txt": "txt", ".md": "md", ".markdown": "md", ".csv": "csv",
    ".html": "html", ".htm": "html", ".pdf": "pdf", ".xlsx": "xlsx",
    ".docx": "docx", ".pptx": "pptx", ".json": "json", ".xml": "xml",
}


def _default_settings() -> Any:
    return SimpleNamespace(
        parser_timeout_seconds=30.0,
        max_pdf_pages=5000,
        max_csv_rows=200000,
        max_xlsx_rows=200000,
    )


def _resolve_fmt(token: str | None, file_path: Path | str) -> str:
    if token and token in _KNOWN_FMTS:
        return token
    if token:
        ext = Path(token).suffix.lower()
        if ext in _EXT_TO_FMT:
            return _EXT_TO_FMT[ext]
    ext = Path(file_path).suffix.lower()
    return _EXT_TO_FMT.get(ext, ext.lstrip(".") or "unknown")


# ---------------------------------------------------------------------------
# ParseResult
# ---------------------------------------------------------------------------

@dataclass
class ExtractedSection:
    heading: str | None
    text: str
    page: int | None = None
    is_table: bool = False


@dataclass
class ParseResult:
    extracted_text: str = ""
    sections: list[ExtractedSection] = field(default_factory=list)
    page_count: int = 0
    char_count: int = 0
    extraction_quality: str = "empty"
    error_message: str | None = None
    format_used: str = ""

    # Aliases expected by tests / callers
    @property
    def error(self) -> str | None:
        return self.error_message

    @property
    def quality(self) -> str:
        return self.extraction_quality

    @classmethod
    def from_error(cls, error: str, fmt: str = "") -> "ParseResult":
        return cls(extraction_quality="failed", error_message=error, format_used=fmt)

    @classmethod
    def from_sections(cls, sections: list[ExtractedSection], page_count: int, fmt: str) -> "ParseResult":
        full_text = "\n\n".join(
            (f"[{s.heading}]\n" if s.heading else "") + s.text
            for s in sections if s.text.strip()
        ).strip()
        char_count = len(full_text)
        if char_count == 0:
            quality = "empty"
        elif char_count < 50:
            quality = "partial"
        else:
            quality = "ok"
        return cls(
            extracted_text=full_text,
            sections=sections,
            page_count=page_count,
            char_count=char_count,
            extraction_quality=quality,
            format_used=fmt,
        )


# ---------------------------------------------------------------------------
# Subprocess worker
# ---------------------------------------------------------------------------

def _subprocess_worker(file_path, fmt, max_pdf_pages, max_csv_rows, max_xlsx_rows,
                       result_queue, error_queue) -> None:
    try:
        path = Path(file_path)
        if fmt == "pdf":
            result = _parse_pdf(path, max_pdf_pages)
        elif fmt in ("txt", "md", "json", "xml"):
            result = _parse_text(path, fmt)
        elif fmt == "html":
            result = _parse_html(path)
        elif fmt == "csv":
            result = _parse_csv(path, max_csv_rows)
        elif fmt == "xlsx":
            result = _parse_xlsx(path, max_xlsx_rows)
        elif fmt == "docx":
            result = _parse_docx(path)
        elif fmt == "pptx":
            result = _parse_pptx(path)
        else:
            result = ParseResult.from_error(f"Unknown format: {fmt!r}", fmt=fmt)
        result_queue.put(result)
    except Exception as exc:
        error_queue.put(f"{type(exc).__name__}: {exc}\n{traceback.format_exc()}"[:4000])


def run_parser_subprocess(
    file_path: Path | str,
    fmt: str | None = None,
    settings: Any | None = None,
) -> ParseResult:
    """Parse a file in an isolated subprocess with a hard timeout. Never raises."""
    settings = settings or _default_settings()
    resolved = _resolve_fmt(fmt, file_path)
    path = str(Path(file_path).resolve())
    timeout = getattr(settings, "parser_timeout_seconds", 30.0)

    try:
        ctx = multiprocessing.get_context("spawn")
    except ValueError:
        ctx = multiprocessing  # type: ignore[assignment]

    result_queue = ctx.Queue()
    error_queue = ctx.Queue()

    proc = ctx.Process(
        target=_subprocess_worker,
        args=(
            path, resolved,
            getattr(settings, "max_pdf_pages", 5000),
            getattr(settings, "max_csv_rows", 200000),
            getattr(settings, "max_xlsx_rows", 200000),
            result_queue, error_queue,
        ),
        daemon=True,
    )

    import queue as _queue
    import time as _time

    proc.start()
    # The result must be READ while waiting: a child that has put a large
    # result on the queue cannot exit until the pipe is drained, so joining
    # first deadlocks every non-trivial document until the timeout kills it.
    deadline = _time.monotonic() + timeout
    result: ParseResult | None = None
    while result is None:
        remaining = deadline - _time.monotonic()
        if remaining <= 0:
            break
        try:
            result = result_queue.get(timeout=min(0.5, remaining))
        except _queue.Empty:
            if not proc.is_alive():
                break  # exited without a result (error path below)

    if result is None and proc.is_alive():
        proc.kill()
        proc.join()
        msg = f"Parser subprocess killed after {timeout:.0f}s timeout (format={resolved})"
        logger.warning(msg)
        return ParseResult.from_error(msg, fmt=resolved)

    proc.join(timeout=10)
    if proc.is_alive():  # result received but the child will not exit
        proc.kill()
        proc.join()
    if result is not None:
        return result

    try:
        err_msg = error_queue.get(timeout=1)
    except Exception:
        err_msg = f"Parser subprocess exited with code {proc.exitcode} and no result"
    return ParseResult.from_error(err_msg[:2000], fmt=resolved)


# ---------------------------------------------------------------------------
# Format parsers (run inside subprocess)
# ---------------------------------------------------------------------------

def _parse_pdf(path: Path, max_pages: int) -> ParseResult:
    import pypdf
    sections: list[ExtractedSection] = []
    pages_processed = 0
    try:
        reader = pypdf.PdfReader(str(path))
        for i in range(min(len(reader.pages), max_pages)):
            try:
                text = (reader.pages[i].extract_text() or "").strip()
                if text:
                    sections.append(ExtractedSection(heading=None, text=text, page=i + 1))
            except Exception as exc:
                logger.debug("PDF page %d error: %s", i + 1, exc)
            pages_processed += 1
    except Exception as exc:
        return ParseResult.from_error(f"PDF parsing failed: {type(exc).__name__}: {exc}", fmt="pdf")
    return ParseResult.from_sections(sections, pages_processed, "pdf")


def _parse_text(path: Path, fmt: str) -> ParseResult:
    try:
        import chardet
        with path.open("rb") as fh:
            sample = fh.read(32768)
        encoding = (chardet.detect(sample).get("encoding") if sample else None) or "utf-8"
    except Exception:
        encoding = "utf-8"

    sections: list[ExtractedSection] = []
    current: list[str] = []
    line_count = 0
    try:
        with path.open("r", encoding=encoding, errors="replace") as fh:
            for raw in fh:
                line = raw.rstrip("\r\n")
                line_count += 1
                if line.strip():
                    current.append(line)
                elif current:
                    para = " ".join(current).strip()
                    if para:
                        sections.append(ExtractedSection(heading=None, text=para, page=None))
                    current = []
            if current:
                para = " ".join(current).strip()
                if para:
                    sections.append(ExtractedSection(heading=None, text=para, page=None))
    except Exception as exc:
        return ParseResult.from_error(f"{fmt.upper()} parsing failed: {exc}", fmt=fmt)
    return ParseResult.from_sections(sections, line_count, fmt)


def _parse_html(path: Path) -> ParseResult:
    from lxml import etree
    TEXT_TAGS = frozenset(["p", "li", "dt", "dd", "caption", "figcaption", "blockquote", "pre"])
    HEADING_TAGS = frozenset(["h1", "h2", "h3", "h4", "h5", "h6"])
    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    table_rows: list[str] = []
    table_cells: list[str] = []
    in_table = in_tr = False
    elements = 0
    try:
        with path.open("rb") as fh:
            for event, elem in etree.iterparse(fh, events=("start", "end"), html=True):
                tag = elem.tag if isinstance(elem.tag, str) else ""
                if event == "start":
                    if tag == "table":
                        in_table, table_rows = True, []
                    elif tag == "tr" and in_table:
                        in_tr, table_cells = True, []
                else:
                    if tag in ("td", "th") and in_table and in_tr:
                        table_cells.append("".join(elem.itertext()).strip())
                    elif tag == "tr" and in_table:
                        if any(c for c in table_cells):
                            table_rows.append(" | ".join(table_cells))
                        in_tr, table_cells = False, []
                    elif tag == "table":
                        if table_rows:
                            sections.append(ExtractedSection(heading=current_heading,
                                                             text="\n".join(table_rows), is_table=True))
                        in_table = False
                    elif tag in HEADING_TAGS and not in_table:
                        t = "".join(elem.itertext()).strip()
                        if t:
                            current_heading = t
                    elif tag in TEXT_TAGS and not in_table:
                        t = "".join(elem.itertext()).strip()
                        if t:
                            sections.append(ExtractedSection(heading=current_heading, text=t))
                    elements += 1
                    elem.clear()
    except Exception as exc:
        return ParseResult.from_error(f"HTML parsing failed: {exc}", fmt="html")
    return ParseResult.from_sections(sections, elements, "html")


def _parse_csv(path: Path, max_rows: int) -> ParseResult:
    try:
        import chardet
        with path.open("rb") as fh:
            sample = fh.read(32768)
        encoding = (chardet.detect(sample).get("encoding") if sample else None) or "utf-8"
    except Exception:
        encoding = "utf-8"

    sections: list[ExtractedSection] = []
    rows_processed = 0
    try:
        with path.open("r", encoding=encoding, errors="replace", newline="") as fh:
            reader = csv.reader(fh)
            header: list[str] | None = None
            batch: list[str] = []
            for row in reader:
                if rows_processed >= max_rows:
                    break
                if not any(c.strip() for c in row):
                    continue
                if header is None:
                    header = [c.strip() for c in row]
                    continue
                pairs = [
                    f"{header[i] if i < len(header) else f'col{i}'}: {c.strip()}"
                    for i, c in enumerate(row) if c.strip()
                ]
                batch.append("; ".join(pairs))
                rows_processed += 1
                if len(batch) >= 50:
                    sections.append(ExtractedSection(heading="Records (CSV)", text="\n".join(batch)))
                    batch = []
            if batch:
                sections.append(ExtractedSection(heading="Records (CSV)", text="\n".join(batch)))
    except Exception as exc:
        return ParseResult.from_error(f"CSV parsing failed: {exc}", fmt="csv")
    return ParseResult.from_sections(sections, rows_processed, "csv")


def _parse_xlsx(path: Path, max_rows: int) -> ParseResult:
    import openpyxl
    sections: list[ExtractedSection] = []
    total_rows = 0
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            for sheet in wb.worksheets:
                sheet_rows: list[str] = []
                header: list[str] | None = None
                for row in sheet.rows:
                    if total_rows >= max_rows:
                        break
                    cells = [str(c.value).strip() if c.value is not None else "" for c in row]
                    if not any(cells):
                        continue
                    if header is None:
                        header = cells
                        continue
                    pairs = [
                        f"{header[i] if i < len(header) else f'col{i}'}: {cell}"
                        for i, cell in enumerate(cells) if cell.strip()
                    ]
                    sheet_rows.append("; ".join(pairs))
                    total_rows += 1
                if sheet_rows:
                    sections.append(ExtractedSection(heading=f"Sheet: {sheet.title}",
                                                     text="\n".join(sheet_rows)))
        finally:
            wb.close()
    except Exception as exc:
        return ParseResult.from_error(f"XLSX parsing failed: {exc}", fmt="xlsx")
    return ParseResult.from_sections(sections, total_rows, "xlsx")


def _parse_docx(path: Path) -> ParseResult:
    from docx import Document
    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    para_count = 0
    try:
        doc = Document(str(path))
        for child in doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                from docx.text.paragraph import Paragraph
                para = Paragraph(child, doc)
                text = para.text.strip()
                if not text:
                    continue
                style = para.style.name if para.style else ""
                if style.startswith("Heading"):
                    current_heading = text
                else:
                    sections.append(ExtractedSection(heading=current_heading, text=text))
                para_count += 1
            elif tag == "tbl":
                from docx.table import Table
                rows = []
                for row in Table(child, doc).rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(c for c in cells):
                        rows.append(" | ".join(cells))
                if rows:
                    sections.append(ExtractedSection(heading=current_heading,
                                                     text="\n".join(rows), is_table=True))
    except Exception as exc:
        return ParseResult.from_error(f"DOCX parsing failed: {exc}", fmt="docx")
    return ParseResult.from_sections(sections, para_count, "docx")


def _parse_pptx(path: Path) -> ParseResult:
    from pptx import Presentation
    sections: list[ExtractedSection] = []
    try:
        prs = Presentation(str(path))
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_title: str | None = None
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if not hasattr(shape, "text_frame"):
                    continue
                is_title = (
                    hasattr(shape, "placeholder_format")
                    and shape.placeholder_format is not None
                    and shape.placeholder_format.idx == 0
                )
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    if is_title and slide_title is None:
                        slide_title = text
                    else:
                        slide_texts.append(text)
            if slide_title or slide_texts:
                heading = f"Slide {slide_num}: {slide_title}" if slide_title else f"Slide {slide_num}"
                sections.append(ExtractedSection(heading=heading, text="\n".join(slide_texts), page=slide_num))
    except Exception as exc:
        return ParseResult.from_error(f"PPTX parsing failed: {exc}", fmt="pptx")
    return ParseResult.from_sections(sections, len(sections), "pptx")


__all__ = ["ParseResult", "ExtractedSection", "run_parser_subprocess"]